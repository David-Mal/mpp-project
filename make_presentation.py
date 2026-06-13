"""
Estethis — Feature Presentation Generator
Produces a dark/gold PowerPoint teaser for the live demo.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

# ── Palette ───────────────────────────────────────────────────
BG       = RGBColor(0x0c, 0x0c, 0x0c)
GOLD     = RGBColor(0xc9, 0xa8, 0x4c)
GOLD_LT  = RGBColor(0xe8, 0xc8, 0x78)
GOLD_DIM = RGBColor(0x6a, 0x5a, 0x28)
TEXT     = RGBColor(0xf0, 0xeb, 0xe0)
TEXT_DIM = RGBColor(0x78, 0x70, 0x58)
TEAL     = RGBColor(0x6a, 0xac, 0x8a)
RED      = RGBColor(0xe0, 0x6c, 0x75)
CARD_BG  = RGBColor(0x16, 0x14, 0x0c)

# ── Helpers ───────────────────────────────────────────────────

def rect(slide, l, t, w, h, fill=None, line=None, lw=Pt(0.75)):
    from pptx.util import Emu
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    s = slide.shapes.add_shape(1,
        Inches(l), Inches(t), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = line; s.line.width = lw
    else:
        s.line.fill.background()
    return s

def tbx(slide, text, l, t, w, h,
        font="Georgia", size=24, color=TEXT,
        bold=False, italic=False, align=PP_ALIGN.LEFT,
        wrap=True):
    box = slide.shapes.add_textbox(
        Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box

def gold_line(slide, l, t, w, h=0.022):
    rect(slide, l, t, w, h, fill=GOLD)

def slide_chrome(slide, label, title_lines, title_size=46, hero=False):
    """Shared layout: left gold bar + label + title + divider."""
    # Background
    bg = slide.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG
    # Left accent bar
    bar_w = 0.10 if not hero else 0.14
    rect(slide, 0, 0, bar_w, 7.5, fill=GOLD)
    # Slide label
    lx = bar_w + 0.55
    tbx(slide, label, lx, 0.32, 6, 0.38,
        font="Calibri Light", size=9, color=GOLD_DIM)
    # Title
    tbx(slide, title_lines, lx, 0.72, 10.8, 2.2,
        font="Georgia", size=title_size, color=GOLD if hero else TEXT)
    # Divider
    gold_line(slide, lx, 2.94, 5.4)
    return lx   # return left-x so callers can align bullets

def bullets(slide, items, lx, top, w=11.2, h=3.8,
            size=17, color=TEXT, dot_color=GOLD):
    box = slide.shapes.add_textbox(
        Inches(lx), Inches(top), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(7)
        dot = p.add_run()
        dot.text = "✦  "
        dot.font.name = "Calibri Light"
        dot.font.size = Pt(size - 2)
        dot.font.color.rgb = dot_color
        body = p.add_run()
        body.text = item
        body.font.name = "Calibri Light"
        body.font.size = Pt(size)
        body.font.color.rgb = color

def stat_box(slide, value, label, l, t, w=1.8, h=1.15,
             val_color=GOLD, lbl_color=TEXT_DIM):
    rect(slide, l, t, w, h, fill=CARD_BG, line=GOLD, lw=Pt(0.6))
    tbx(slide, value, l + 0.12, t + 0.08, w - 0.24, 0.55,
        font="Georgia", size=28, color=val_color, align=PP_ALIGN.CENTER)
    tbx(slide, label, l + 0.12, t + 0.62, w - 0.24, 0.45,
        font="Calibri Light", size=10, color=lbl_color, align=PP_ALIGN.CENTER)

# ── Presentation setup ────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# ═══════════════════════════════════════════════════════════════
# S01 — TITLE
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank)
bg = sl.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG
rect(sl, 0, 0, 0.10, 7.5, fill=GOLD)

# Ornament
tbx(sl, "✦", 1.0, 1.25, 1.0, 0.9,
    font="Georgia", size=34, color=GOLD, align=PP_ALIGN.LEFT)

# Brand name — huge
tbx(sl, "ESTETHIS", 1.65, 1.0, 11, 1.8,
    font="Georgia", size=96, color=TEXT, bold=False)

# Tagline
rect(sl, 1.65, 3.0, 6.5, 0.022, fill=GOLD)
tbx(sl, "Luxury Fashion E-Commerce Platform", 1.65, 3.15, 10, 0.7,
    font="Calibri Light", size=22, color=GOLD)

# Tech stack
tbx(sl,
    "React 19  ·  Vite  ·  Express 5  ·  Sequelize  ·  WebSocket  ·  JWT Auth",
    1.65, 3.9, 11, 0.55,
    font="Calibri Light", size=13, color=TEXT_DIM)

tbx(sl, "8 Feature Phases  ·  40+ Components  ·  Full-Stack",
    1.65, 4.45, 11, 0.5,
    font="Calibri Light", size=13, color=GOLD_DIM)

tbx(sl, "EST. 2026", 1.65, 6.6, 4, 0.5,
    font="Calibri Light", size=10, color=GOLD_DIM)

# ═══════════════════════════════════════════════════════════════
# S02 — ROLE-BASED ACCESS CONTROL
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank)
lx = slide_chrome(sl, "01  ·  SECURITY", "Not All\nUsers Are Equal")

bullets(sl, [
    "3-tier permission model: admin  /  manager  /  user",
    "Every UI element gates in real-time — no page reload, no flash of restricted content",
    "Backend enforces each route independently (UI gating is UX courtesy, not the security layer)",
    "Managers edit the catalogue · Admins access user management, action logs & threat panel",
], lx, 3.1)

# Role cards on the right
for i, (role, col) in enumerate([("ADMIN", GOLD), ("MANAGER", TEAL), ("USER", TEXT_DIM)]):
    stat_box(sl, role, ["Full access", "Catalogue", "Browse"][i],
             10.5, 2.5 + i * 1.42, w=2.1, h=1.2,
             val_color=col, lbl_color=TEXT_DIM)

# ═══════════════════════════════════════════════════════════════
# S03 — GUEST MODE
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank)
lx = slide_chrome(sl, "02  ·  AUTH FLOW", "Browse Before\nYou Commit")

bullets(sl, [
    "Full product catalogue — including reviews — accessible without an account",
    "Protected actions (Add to Cart, Checkout, My Account) trigger an intercept modal, not a redirect",
    "Auth state machine:  init  →  guest  →  login  /  register  →  app",
    "View state is preserved across the auth round-trip — user lands back exactly where they were",
], lx, 3.1)

# State machine diagram (simple boxes + arrows)
for i, (stage, col) in enumerate([
    ("GUEST", TEXT_DIM), ("INTERCEPT", GOLD), ("LOGIN", TEAL), ("APP", TEXT)
]):
    stat_box(sl, stage, ["Browse", "Modal", "Auth", "Full access"][i],
             0.72 + i * 2.82, 5.55, w=2.42, h=1.12,
             val_color=col, lbl_color=TEXT_DIM)
    if i < 3:
        tbx(sl, "→", 3.0 + i * 2.82, 5.9, 0.6, 0.5,
            font="Georgia", size=18, color=GOLD_DIM, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════════════
# S04 — THE TICKET MACHINE  (hero slide)
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank)
lx = slide_chrome(sl, "03  ·  GAMIFICATION", "The Ticket\nMachine",
                  title_size=54, hero=True)

# Dramatic reel symbols watermark
tbx(sl, "✦ ✦ ✦", 7.8, 1.6, 5.5, 2.5,
    font="Georgia", size=72, color=CARD_BG, align=PP_ALIGN.CENTER)

# Win-chance callout
stat_box(sl, "5 %", "WIN CHANCE", 10.2, 1.0, w=2.4, h=1.3, val_color=GOLD)
stat_box(sl, "3", "PRIZE TYPES", 10.2, 2.52, w=2.4, h=1.2, val_color=GOLD_LT)
stat_box(sl, "30d", "COUPON TTL", 10.2, 3.94, w=2.4, h=1.2, val_color=TEAL)

bullets(sl, [
    "Result computed once at mount — reels reveal a pre-determined outcome",
    "CSS frame-by-frame animation: 65 ms/frame, 3 reels stop sequentially 400 ms apart",
    "Prize types: Percentage OFF  ·  Fixed $ OFF  ·  Buy 1 Get 1",
    "Won coupons stored in user-scoped localStorage with 30-day expiry",
    "Redeemable from My Account → My Coupons with one-click copy-to-clipboard",
], lx, 3.12, w=9.5, size=16)

# ═══════════════════════════════════════════════════════════════
# S05 — SMART PERSISTENCE
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank)
lx = slide_chrome(sl, "04  ·  PERSISTENCE", "Your Session\nSurvives Everything")

bullets(sl, [
    "User-scoped cart key: estethis_cart_u{id} — cart never leaks across accounts on the same browser",
    "Page refresh: apiMe() validates token and reloads the user's cart in a single pass",
    "Cookie activity tracker records recently-viewed categories across browser sessions",
    "Personalization banner in the catalogue surfaces quick-filter chips from cookie history",
    "Inactivity auto-logout (30 min) saves the cart to the scoped key before clearing",
], lx, 3.1)

# Storage layer diagram
for i, (label, sub, col) in enumerate([
    ("localStorage", "cart · orders · coupons", GOLD),
    ("Cookies", "categories · last page", TEAL),
    ("sessionStorage", "UI state (banner dismiss)", TEXT_DIM),
]):
    stat_box(sl, label, sub, lx + i * 3.72, 5.62, w=3.5, h=1.12,
             val_color=col, lbl_color=TEXT_DIM)

# ═══════════════════════════════════════════════════════════════
# S06 — ATELIER CONFIGURATOR
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank)
lx = slide_chrome(sl, "05  ·  BESPOKE TAILORING", "The Atelier\nConfigurator")

bullets(sl, [
    "Interactive SVG body silhouette — 5 annotated measurement points, click to focus",
    "3-step wizard: Select Product  →  Mark & Measure  →  Confirm & Send",
    "Strict per-field validation: required · positive · realistic range  (e.g. chest 60–200 cm)",
    "Errors clear in real-time as the user types — error state never bleeds into Step 3",
    "Rush delivery + premium packaging add-ons update the pricing summary live",
], lx, 3.1)

# Measurement range cards
for i, (field, rng) in enumerate([
    ("Shoulder", "25–70 cm"), ("Chest", "60–200 cm"),
    ("Sleeve", "40–90 cm"), ("Waist", "50–200 cm"), ("Length", "40–130 cm"),
]):
    stat_box(sl, field, rng, lx + i * 2.32, 5.62, w=2.18, h=1.1,
             val_color=GOLD if i == 0 else TEXT_DIM, lbl_color=TEXT_DIM)

# ═══════════════════════════════════════════════════════════════
# S07 — ADMIN COMMAND CENTER
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank)
lx = slide_chrome(sl, "06  ·  ADMIN PANEL", "Admin\nCommand Center")

bullets(sl, [
    "Orders tab: every customer order, fully expandable with items, prices & shipping address",
    "Inbox tab: contact-form messages with unread badge · mark-as-read on expand",
    "Sales KPIs injected into Statistics: Total Revenue + Orders Placed alongside product KPIs",
    "Role-gated: Inbox is admin-only — admins see a panel notice on the Contact page, not the form",
], lx, 3.1)

# Tab cards
for i, (tab, desc, col) in enumerate([
    ("STATISTICS", "Products + Sales KPIs", GOLD),
    ("ORDERS", "All customer orders", TEAL),
    ("INBOX", "Contact messages", RED),
]):
    stat_box(sl, tab, desc, lx + i * 3.78, 5.6, w=3.55, h=1.15,
             val_color=col, lbl_color=TEXT_DIM)

# ═══════════════════════════════════════════════════════════════
# S08 — REAL-TIME & OFFLINE RESILIENCE
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank)
lx = slide_chrome(sl, "07  ·  REAL-TIME", "Alive Even\nOffline")

bullets(sl, [
    "WebSocket connection — product mutations propagate to all open tabs in real-time",
    "Offline mutation queue — CRUD ops are batched locally and replayed on reconnect",
    "Optimistic UI — changes are applied instantly; rolled back silently on server failure",
    "Three-signal connection indicator: browser online · WebSocket · server health",
], lx, 3.1)

# Signal cards
for i, (sig, status, col) in enumerate([
    ("BROWSER", "navigator.onLine", TEAL),
    ("WEBSOCKET", "ws:// heartbeat", GOLD),
    ("SERVER", "/api/health probe", TEXT_DIM),
]):
    stat_box(sl, sig, status, lx + i * 3.78, 5.6, w=3.55, h=1.15,
             val_color=col, lbl_color=TEXT_DIM)

# ═══════════════════════════════════════════════════════════════
# S09 — DEMO TEASER (closing)
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank)
bg = sl.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG

# Gold border on all four sides
rect(sl, 0, 0, 13.33, 0.08, fill=GOLD)
rect(sl, 0, 7.42, 13.33, 0.08, fill=GOLD)
rect(sl, 0, 0, 0.08, 7.5, fill=GOLD)
rect(sl, 13.25, 0, 0.08, 7.5, fill=GOLD)

# Corner ornaments
for cx, cy in [(0.45, 0.25), (12.0, 0.25), (0.45, 6.8), (12.0, 6.8)]:
    tbx(sl, "✦", cx, cy, 0.8, 0.7,
        font="Georgia", size=14, color=GOLD_DIM, align=PP_ALIGN.CENTER)

# Main text
tbx(sl, "✦", 6.3, 0.95, 0.9, 0.8,
    font="Georgia", size=26, color=GOLD, align=PP_ALIGN.CENTER)

tbx(sl, "Ready to see it\nin action?", 1.5, 1.6, 10.5, 2.8,
    font="Georgia", size=60, color=TEXT, align=PP_ALIGN.CENTER)

rect(sl, 4.4, 4.52, 4.7, 0.028, fill=GOLD)

tbx(sl,
    "Guest Browse  ·  Cart & Checkout  ·  The Ticket Machine\n"
    "My Coupons  ·  Atelier Configurator  ·  Admin Inbox",
    1.5, 4.7, 10.5, 1.1,
    font="Calibri Light", size=14, color=TEXT_DIM, align=PP_ALIGN.CENTER)

# CTA button simulation
rect(sl, 5.15, 5.95, 3.05, 0.7, line=GOLD, lw=Pt(1.2))
tbx(sl, "LIVE DEMO  →", 5.15, 5.97, 3.05, 0.66,
    font="Calibri Light", size=17, color=GOLD, bold=True, align=PP_ALIGN.CENTER)

# ── Save ─────────────────────────────────────────────────────
output = r"C:\Users\David\Desktop\mpp-project\Estethis_Presentation.pptx"
prs.save(output)
print(f"Saved: {output}")
